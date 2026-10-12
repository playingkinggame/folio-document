"""Read PDFs, Word files and plain text, then split them into chunks."""
from dataclasses import dataclass
from pathlib import Path

from langchain_text_splitters import RecursiveCharacterTextSplitter

from .config import settings

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class Page:
    number: int | None  # 1-based page number for PDFs, None for other formats
    text: str


@dataclass
class Chunk:
    text: str
    page: int | None
    index: int


def _read_pdf(path: Path) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(Page(i, text))
    if sum(len(p.text) for p in pages) < 50:
        raise ValueError(
            "No readable text found. This looks like a scanned PDF (images only); run OCR on it first."
        )
    return pages


def _read_docx(path: Path) -> list[Page]:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("The Word file has no text.")
    return [Page(None, text)]


def _read_text(path: Path) -> list[Page]:
    text = path.read_text(encoding="utf-8", errors="ignore").strip()
    if not text:
        raise ValueError("The file is empty.")
    return [Page(None, text)]


def load_pages(path: Path) -> list[Page]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext in (".txt", ".md"):
        return _read_text(path)
    raise ValueError(f"Unsupported file type: {ext}")


def split_pages(pages: list[Page]) -> list[Chunk]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks: list[Chunk] = []
    for page in pages:
        for piece in splitter.split_text(page.text):
            piece = piece.strip()
            if len(piece) >= 20:
                chunks.append(Chunk(piece, page.number, len(chunks)))
    return chunks
